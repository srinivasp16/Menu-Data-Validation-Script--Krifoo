"""
Krifoo Menu Data Validation Script
====================================
Accepts restaurant menu files in any common format, converts them to
a clean CSV, then validates every row before platform upload.

Supported input formats:
  .csv  .tsv  .xlsx  .xls  .ods  .docx  .doc (via docx)  .pdf  .txt  .json

Usage:
    python krifoo_menu_validator.py menu.xlsx
    python krifoo_menu_validator.py menu.pdf --output clean_menu.csv
    python krifoo_menu_validator.py menu.docx --report report.txt --no-colour
"""

import csv
import sys
import re
import json
import argparse
import io
from collections import defaultdict, Counter
from pathlib import Path


# ──────────────────────────────────────────────
# Configuration
# ──────────────────────────────────────────────

REQUIRED_COLUMNS  = {"item_name", "category", "price", "description"}
OPTIONAL_COLUMNS  = {"item_id", "currency", "availability", "tags"}
PRICE_PATTERN     = re.compile(r"^\d+(\.\d{1,2})?$")
ITEM_NAME_PATTERN = re.compile(r"^[\w\s\-&'(),./]+$")
MAX_NAME_LENGTH   = 100
MAX_DESC_LENGTH   = 500

# Column aliases — maps common variations to canonical names
COLUMN_ALIASES: dict[str, str] = {
    "name":            "item_name",
    "item":            "item_name",
    "dish":            "item_name",
    "dish_name":       "item_name",
    "food_name":       "item_name",
    "menu_item":       "item_name",
    "product_name":    "item_name",
    "title":           "item_name",
    "cat":             "category",
    "section":         "category",
    "type":            "category",
    "menu_section":    "category",
    "course":          "category",
    "cost":            "price",
    "rate":            "price",
    "amount":          "price",
    "unit_price":      "price",
    "selling_price":   "price",
    "mrp":             "price",
    "desc":            "description",
    "details":         "description",
    "info":            "description",
    "about":           "description",
    "notes":           "description",
    "avail":           "availability",
    "status":          "availability",
    "in_stock":        "availability",
    "available":       "availability",
}

SUPPORTED_EXTENSIONS = {".csv", ".tsv", ".xlsx", ".xls", ".ods",
                        ".doc", ".docx", ".pdf", ".txt", ".json"}


# ──────────────────────────────────────────────
# Validation report
# ──────────────────────────────────────────────

class ValidationReport:
    def __init__(self, filepath: str):
        self.filepath  = filepath
        self.issues    = []
        self.warnings  = []
        self.info      = []

    def error(self, row: int, code: str, message: str):
        self.issues.append((row, "ERROR", code, message))

    def warn(self, row: int, code: str, message: str):
        self.warnings.append((row, "WARN", code, message))

    def note(self, message: str):
        self.info.append(message)

    @property
    def has_errors(self):
        return bool(self.issues)

    def all_findings(self):
        combined = self.issues + self.warnings
        combined.sort(key=lambda x: x[0])
        return combined

    def summary(self):
        return {"errors": len(self.issues), "warnings": len(self.warnings)}


# ──────────────────────────────────────────────
# Converters  →  list[dict]
# ──────────────────────────────────────────────

def _normalise_header(raw: str) -> str:
    """Lowercase, strip, spaces→underscore, then apply alias map."""
    key = raw.strip().lower().replace(" ", "_")
    return COLUMN_ALIASES.get(key, key)


def _rows_from_records(records: list[dict]) -> list[dict]:
    """Re-key every record through the alias map."""
    out = []
    for rec in records:
        out.append({_normalise_header(k): (str(v).strip() if v is not None else "")
                    for k, v in rec.items()})
    return out


# ── CSV / TSV ──────────────────────────────────
def convert_csv(path: Path, report: ValidationReport) -> list[dict]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    try:
        with open(path, newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f, delimiter=delimiter)
            if reader.fieldnames is None:
                report.error(0, "EMPTY_FILE", "File is empty or has no header.")
                return []
            return _rows_from_records(list(reader))
    except UnicodeDecodeError:
        # Try latin-1 fallback
        with open(path, newline="", encoding="latin-1") as f:
            reader = csv.DictReader(f, delimiter=delimiter)
            report.warn(0, "ENCODING_FALLBACK",
                        "File is not UTF-8; read with latin-1 encoding. "
                        "Re-save as UTF-8 for best results.")
            return _rows_from_records(list(reader))


# ── Excel / ODS ───────────────────────────────
def convert_excel(path: Path, report: ValidationReport) -> list[dict]:
    try:
        import pandas as pd
    except ImportError:
        report.error(0, "MISSING_LIBRARY",
                     "pandas is required for Excel/ODS files. "
                     "Run: pip install pandas openpyxl")
        return []

    try:
        engine = "odf" if path.suffix.lower() == ".ods" else None
        xl = pd.ExcelFile(path, engine=engine)
        sheet_name = xl.sheet_names[0]
        if len(xl.sheet_names) > 1:
            report.warn(0, "MULTIPLE_SHEETS",
                        f"Workbook has {len(xl.sheet_names)} sheets. "
                        f"Using first sheet: '{sheet_name}'.")
        df = xl.parse(sheet_name, dtype=str).fillna("")
        df.columns = [_normalise_header(c) for c in df.columns]
        return df.to_dict(orient="records")
    except Exception as e:
        report.error(0, "EXCEL_PARSE_ERROR", f"Could not read Excel file: {e}")
        return []


# ── DOC (legacy Word) ─────────────────────────
def convert_doc(path: Path, report: ValidationReport) -> list[dict]:
    """Convert old-format .doc to .docx via LibreOffice, then parse."""
    import shutil, subprocess, tempfile

    lo = shutil.which("libreoffice") or shutil.which("soffice")
    if not lo:
        report.error(0, "MISSING_TOOL",
                     "LibreOffice is required to read .doc files. "
                     "Install it from https://www.libreoffice.org or "
                     "convert your file to .docx and try again.")
        return []

    with tempfile.TemporaryDirectory() as tmp:
        try:
            result = subprocess.run(
                [lo, "--headless", "--convert-to", "docx",
                 str(path), "--outdir", tmp],
                capture_output=True, text=True, timeout=60,
            )
        except subprocess.TimeoutExpired:
            report.error(0, "DOC_CONVERT_TIMEOUT",
                         "LibreOffice took too long to convert the .doc file.")
            return []

        converted = Path(tmp) / (path.stem + ".docx")
        if not converted.exists():
            report.error(0, "DOC_CONVERT_FAILED",
                         f"Could not convert '{path.name}' to docx. "
                         f"LibreOffice output: {result.stderr.strip()[:200]}")
            return []

        report.note(f"Converted '{path.name}' → .docx via LibreOffice.")
        return _parse_docx_file(converted, report)


# ── DOCX ──────────────────────────────────────
def convert_docx(path: Path, report: ValidationReport) -> list[dict]:
    return _parse_docx_file(path, report)


def _parse_docx_file(path: Path, report: ValidationReport) -> list[dict]:
    """Core DOCX reader — shared by both .doc and .docx converters."""
    try:
        from docx import Document
    except ImportError:
        report.error(0, "MISSING_LIBRARY",
                     "python-docx is required for Word files. "
                     "Run: pip install python-docx")
        return []

    try:
        doc = Document(str(path))
    except Exception as e:
        report.error(0, "DOCX_READ_ERROR", f"Could not open Word file: {e}")
        return []

    rows: list[dict] = []

    # 1. Try tables first — most reliable structure
    for table in doc.tables:
        if len(table.rows) < 2:
            continue
        headers = [_normalise_header(cell.text) for cell in table.rows[0].cells]
        if not any(h in REQUIRED_COLUMNS for h in headers):
            # Header row doesn't look like a menu — try next table
            continue
        for tr in table.rows[1:]:
            rec = {h: cell.text.strip() for h, cell in zip(headers, tr.cells)}
            if any(v for v in rec.values()):
                rows.append(rec)
        if rows:
            report.note(f"Extracted {len(rows)} rows from Word table.")
            return rows

    # 2. Fall back: try all tables ignoring header check
    for table in doc.tables:
        if len(table.rows) < 2:
            continue
        headers = [_normalise_header(cell.text) for cell in table.rows[0].cells]
        for tr in table.rows[1:]:
            rec = {h: cell.text.strip() for h, cell in zip(headers, tr.cells)}
            if any(v for v in rec.values()):
                rows.append(rec)
        if rows:
            report.note(f"Extracted {len(rows)} rows from Word table (relaxed header check).")
            return rows

    # 3. Fall back to paragraph delimiter detection
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    if not paragraphs:
        report.error(0, "EMPTY_DOCX",
                     "Word document has no readable text or tables.")
        return []

    sep = next((s for s in ["|", "\t", ";", ","] if any(s in p for p in paragraphs)), None)
    if sep:
        headers = [_normalise_header(h) for h in paragraphs[0].split(sep)]
        for line in paragraphs[1:]:
            parts = [p.strip() for p in line.split(sep)]
            if len(parts) >= len(headers):
                rows.append(dict(zip(headers, parts)))
        if rows:
            report.note(f"Extracted {len(rows)} rows from Word paragraphs (separator: '{sep}').")
            return rows

    report.error(0, "DOCX_UNSTRUCTURED",
                 "Could not detect a table or delimited layout in the Word document. "
                 "Please format the menu as a Word table with columns: "
                 "item_name, category, price, description.")
    return []


# ── PDF ───────────────────────────────────────
def convert_pdf(path: Path, report: ValidationReport) -> list[dict]:
    try:
        import pdfplumber
    except ImportError:
        report.error(0, "MISSING_LIBRARY",
                     "pdfplumber is required for PDF files. "
                     "Run: pip install pdfplumber")
        return []

    rows: list[dict] = []

    with pdfplumber.open(str(path)) as pdf:
        report.note(f"PDF has {len(pdf.pages)} page(s).")

        # 1. Try structured tables across all pages
        for page_num, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables()
            for table in tables:
                if not table or len(table) < 2:
                    continue
                headers = [_normalise_header(str(h or "")) for h in table[0]]
                for data_row in table[1:]:
                    rec = {}
                    for h, cell in zip(headers, data_row):
                        rec[h] = str(cell).strip() if cell else ""
                    if any(v for v in rec.values()):
                        rows.append(rec)

        if rows:
            report.note(f"Extracted {len(rows)} rows from PDF tables.")
            return rows

        # 2. Fall back to plain text with separator detection
        full_text = "\n".join(
            page.extract_text() or "" for page in pdf.pages
        )
        lines = [l.strip() for l in full_text.splitlines() if l.strip()]
        if not lines:
            report.error(0, "EMPTY_PDF", "Could not extract any text from the PDF.")
            return []

        sep = None
        for candidate in ["|", "\t", ";"]:
            if any(candidate in l for l in lines):
                sep = candidate
                break

        if sep:
            headers = [_normalise_header(h) for h in lines[0].split(sep)]
            for line in lines[1:]:
                parts = [p.strip() for p in line.split(sep)]
                if len(parts) >= len(headers):
                    rows.append(dict(zip(headers, parts)))
            if rows:
                report.note(f"Extracted {len(rows)} rows from PDF text (separator: '{sep}').")
                return rows

    report.error(0, "PDF_UNSTRUCTURED",
                 "Could not detect a table or delimited layout in the PDF. "
                 "For best results, export the menu as an Excel or CSV file.")
    return []


# ── TXT ───────────────────────────────────────
def convert_txt(path: Path, report: ValidationReport) -> list[dict]:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    if not lines:
        report.error(0, "EMPTY_FILE", "Text file is empty.")
        return []

    sep = None
    for candidate in ["\t", "|", ";", ","]:
        if any(candidate in l for l in lines):
            sep = candidate
            break

    if not sep:
        report.error(0, "TXT_UNSTRUCTURED",
                     "Cannot detect a delimiter (tab, |, ;, ,) in the text file.")
        return []

    headers = [_normalise_header(h) for h in lines[0].split(sep)]
    rows = []
    for line in lines[1:]:
        parts = [p.strip() for p in line.split(sep)]
        if len(parts) >= len(headers):
            rows.append(dict(zip(headers, parts)))
    report.note(f"Extracted {len(rows)} rows from text file (separator: '{sep}').")
    return rows


# ── JSON ──────────────────────────────────────
def convert_json(path: Path, report: ValidationReport) -> list[dict]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        report.error(0, "JSON_PARSE_ERROR", f"Invalid JSON: {e}")
        return []

    # Accept: list of dicts  OR  {"menu": [...]}  OR  {"items": [...]}
    if isinstance(data, list):
        records = data
    elif isinstance(data, dict):
        for key in ("menu", "items", "data", "products"):
            if key in data and isinstance(data[key], list):
                records = data[key]
                break
        else:
            report.error(0, "JSON_STRUCTURE",
                         "Expected a JSON array or an object with a "
                         "'menu', 'items', or 'data' array.")
            return []
    else:
        report.error(0, "JSON_STRUCTURE", "JSON root must be an array or object.")
        return []

    return _rows_from_records(records)


# ──────────────────────────────────────────────
# Format dispatcher
# ──────────────────────────────────────────────

def load_file(path: Path, report: ValidationReport) -> list[dict]:
    ext = path.suffix.lower()
    report.note(f"Input format: {ext.upper().lstrip('.')}")

    if ext in (".csv", ".tsv"):
        return convert_csv(path, report)
    elif ext in (".xlsx", ".xls", ".ods"):
        return convert_excel(path, report)
    elif ext == ".doc":
        return convert_doc(path, report)
    elif ext == ".docx":
        return convert_docx(path, report)
    elif ext == ".pdf":
        return convert_pdf(path, report)
    elif ext == ".txt":
        return convert_txt(path, report)
    elif ext == ".json":
        return convert_json(path, report)
    else:
        report.error(0, "UNSUPPORTED_FORMAT",
                     f"Unsupported file type '{ext}'. "
                     f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")
        return []


# ──────────────────────────────────────────────
# Validators
# ──────────────────────────────────────────────

def check_columns(headers: list[str], report: ValidationReport) -> set[str]:
    normalised = set(headers)
    missing = REQUIRED_COLUMNS - normalised
    for col in sorted(missing):
        report.error(0, "MISSING_COLUMN",
                     f"Required column '{col}' not found. "
                     f"Available columns: {', '.join(sorted(normalised))}")
    extra = normalised - REQUIRED_COLUMNS - OPTIONAL_COLUMNS
    for col in sorted(extra):
        report.warn(0, "UNKNOWN_COLUMN",
                    f"Unrecognised column '{col}' — will be kept as-is.")
    return normalised


def check_missing_price(row: dict, row_num: int, report: ValidationReport):
    price = row.get("price", "").strip().lstrip("₹$€£¥")
    if not price:
        report.error(row_num, "MISSING_PRICE", f"Row {row_num}: 'price' is empty.")
    elif not PRICE_PATTERN.match(price):
        report.error(row_num, "INVALID_PRICE",
                     f"Row {row_num}: price '{price}' is not a valid number "
                     f"(expected: 9.99 or 10).")
    elif float(price) <= 0:
        report.error(row_num, "ZERO_PRICE",
                     f"Row {row_num}: price must be greater than 0.")


def check_empty_category(row: dict, row_num: int, report: ValidationReport):
    if not row.get("category", ""):
        report.error(row_num, "EMPTY_CATEGORY", f"Row {row_num}: 'category' is empty.")


def check_item_name(row: dict, row_num: int, report: ValidationReport) -> str:
    name = row.get("item_name", "")
    if not name:
        report.error(row_num, "MISSING_ITEM_NAME", f"Row {row_num}: 'item_name' is empty.")
        return ""
    if len(name) > MAX_NAME_LENGTH:
        report.warn(row_num, "NAME_TOO_LONG",
                    f"Row {row_num}: item name exceeds {MAX_NAME_LENGTH} chars ({len(name)}).")
    if not ITEM_NAME_PATTERN.match(name):
        report.warn(row_num, "SUSPICIOUS_NAME_CHARS",
                    f"Row {row_num}: item name '{name}' contains unusual characters.")
    return name


def check_description(row: dict, row_num: int, report: ValidationReport):
    desc = row.get("description", "")
    if not desc:
        report.warn(row_num, "EMPTY_DESCRIPTION", f"Row {row_num}: 'description' is empty.")
    elif len(desc) > MAX_DESC_LENGTH:
        report.warn(row_num, "DESC_TOO_LONG",
                    f"Row {row_num}: description exceeds {MAX_DESC_LENGTH} chars.")


def check_availability(row: dict, row_num: int, report: ValidationReport):
    avail = row.get("availability", "")
    valid = {"yes", "no", "true", "false", "1", "0", "available", "unavailable"}
    if avail and avail.lower() not in valid:
        report.warn(row_num, "INVALID_AVAILABILITY",
                    f"Row {row_num}: availability '{avail}' is non-standard "
                    f"(expected: yes/no, true/false, 1/0).")


def check_formatting_consistency(rows: list[dict], report: ValidationReport):
    category_variants: dict[str, set] = defaultdict(set)
    for row in rows:
        cat = row.get("category", "")
        if cat:
            category_variants[cat.lower()].add(cat)
    for canonical, variants in category_variants.items():
        if len(variants) > 1:
            report.warn(0, "INCONSISTENT_CATEGORY_CASE",
                        f"Category '{canonical}' has mixed capitalisation: "
                        + ", ".join(f"'{v}'" for v in sorted(variants)))


def check_duplicate_names(names: list[str], report: ValidationReport):
    counts = Counter(n.lower() for n in names if n)
    for name, count in counts.items():
        if count > 1:
            report.error(0, "DUPLICATE_ITEM_NAME",
                         f"Item name '{name}' appears {count} times. Names must be unique.")


# ──────────────────────────────────────────────
# CSV writer
# ──────────────────────────────────────────────

OUTPUT_COLUMNS = ["item_name", "category", "price", "description",
                  "availability", "item_id", "currency", "tags"]

def write_csv(rows: list[dict], output_path: Path):
    """Write rows to a clean, standardised CSV."""
    # Discover all column keys present
    all_keys: list[str] = []
    seen: set[str] = set()
    for preferred in OUTPUT_COLUMNS:
        if preferred not in seen:
            all_keys.append(preferred)
            seen.add(preferred)
    for row in rows:
        for k in row:
            if k not in seen:
                all_keys.append(k)
                seen.add(k)

    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=all_keys, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            # Strip currency symbols from price before writing
            price = row.get("price", "")
            row["price"] = price.lstrip("₹$€£¥").strip()
            writer.writerow({k: row.get(k, "") for k in all_keys})


# ──────────────────────────────────────────────
# Orchestrator
# ──────────────────────────────────────────────

def validate_menu(filepath: str, output_csv: str | None = None) -> tuple[ValidationReport, Path | None]:
    report = ValidationReport(filepath)
    path   = Path(filepath)

    if not path.exists():
        report.error(0, "FILE_NOT_FOUND", f"File '{filepath}' does not exist.")
        return report, None

    rows = load_file(path, report)
    if not rows:
        return report, None

    # Normalise — strip whitespace from all values
    rows = [{k: v.strip() for k, v in row.items()} for row in rows]

    # Validate columns
    all_cols = list({k for row in rows for k in row})
    available_cols = check_columns(all_cols, report)

    item_names: list[str] = []
    for i, row in enumerate(rows, start=2):
        name = check_item_name(row, i, report)
        item_names.append(name)
        check_missing_price(row, i, report)
        check_empty_category(row, i, report)
        check_description(row, i, report)
        if "availability" in available_cols:
            check_availability(row, i, report)

    check_duplicate_names(item_names, report)
    check_formatting_consistency(rows, report)

    unique_cats = len({r.get("category", "").lower() for r in rows if r.get("category")})
    report.note(f"Total rows checked  : {len(rows)}")
    report.note(f"Unique categories   : {unique_cats}")

    # Write CSV output
    csv_path: Path | None = None
    if output_csv:
        csv_path = Path(output_csv)
    elif path.suffix.lower() != ".csv":
        csv_path = path.with_suffix(".csv")

    if csv_path:
        write_csv(rows, csv_path)
        report.note(f"CSV output saved to : {csv_path}")

    return report, csv_path


# ──────────────────────────────────────────────
# Display
# ──────────────────────────────────────────────

SEVERITY_COLOURS = {
    "ERROR": "\033[91m", "WARN": "\033[93m",
    "INFO":  "\033[96m", "RESET": "\033[0m",
    "BOLD":  "\033[1m",  "GREEN": "\033[92m",
}

def colour(text: str, code: str, use_colour: bool) -> str:
    if not use_colour:
        return text
    return f"{SEVERITY_COLOURS[code]}{text}{SEVERITY_COLOURS['RESET']}"


def print_report(report: ValidationReport, use_colour: bool = True) -> str:
    lines = []
    sep   = "─" * 62
    lines.append(sep)
    lines.append(colour("  KRIFOO MENU VALIDATOR", "BOLD", use_colour))
    lines.append(f"  File : {report.filepath}")
    lines.append(sep)
    for note in report.info:
        lines.append(f"  ℹ  {note}")
    lines.append("")
    findings = report.all_findings()
    if not findings:
        lines.append(colour("  ✅  No issues found — menu is ready to upload!", "GREEN", use_colour))
    else:
        for row, sev, code, msg in findings:
            icon     = "✖" if sev == "ERROR" else "⚠"
            tag      = colour(f"[{sev}]", sev, use_colour)
            location = f"(row {row})" if row > 0 else "(file)"
            lines.append(f"  {icon} {tag} {code} {location}")
            lines.append(f"     {msg}")
    lines.append("")
    lines.append(sep)
    s         = report.summary()
    err_text  = colour(f"{s['errors']} error(s)",   "ERROR", use_colour)
    warn_text = colour(f"{s['warnings']} warning(s)", "WARN", use_colour)
    verdict   = colour("FAILED ✖", "ERROR", use_colour) if report.has_errors \
                else colour("PASSED ✅", "GREEN", use_colour)
    lines.append(f"  Result : {verdict}   {err_text}   {warn_text}")
    lines.append(sep)
    return "\n".join(lines)


# ──────────────────────────────────────────────
# CLI
# ──────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Krifoo — Validate (and convert) a restaurant menu file before upload.",
        epilog=f"Supported formats: {', '.join(sorted(SUPPORTED_EXTENSIONS))}",
    )
    parser.add_argument("menu_file", help="Path to the menu file.")
    parser.add_argument("--output", "-o", metavar="FILE",
                        help="Path for the converted CSV output "
                             "(default: same name as input with .csv extension).")
    parser.add_argument("--report", "-r", metavar="FILE",
                        help="Write a plain-text validation report to this file.")
    parser.add_argument("--no-colour", action="store_true",
                        help="Disable ANSI colour output (useful for CI logs).")
    args = parser.parse_args()

    report, csv_path = validate_menu(args.menu_file, args.output)
    use_colour = not args.no_colour and sys.stdout.isatty()
    output     = print_report(report, use_colour=use_colour)

    print(output)

    if args.report:
        Path(args.report).write_text(
            print_report(report, use_colour=False), encoding="utf-8")
        print(f"\n  Report saved to: {args.report}")

    sys.exit(1 if report.has_errors else 0)


if __name__ == "__main__":
    main()
